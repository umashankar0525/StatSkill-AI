"""Real converter checks; creates only temporary synthetic test material."""
import os
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import rag_engine as rag
from PIL import Image,ImageDraw,ImageFont
with tempfile.TemporaryDirectory() as folder:
    root=Path(folder)
    line='Stratified sampling improves survey precision.'
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',36)
    im=Image.new('RGB',(1300,180),'white');ImageDraw.Draw(im).text((20,45),line,font=font,fill='black')
    for ext in ['png','jpg','tif','bmp','webp']:
        p=root/('scan.'+ext);im.save(p)
        result=rag.extract_text_from_file(p)
        assert 'sampling' in result[0]['text'].lower(),result
        print('OCR PASS',ext,flush=True)
    import fitz
    doc=fitz.open();page=doc.new_page(width=1300,height=180);page.insert_image(page.rect,filename=str(root/'scan.png'));doc.save(root/'scan.pdf');doc.close()
    assert 'sampling' in rag.extract_text_from_file(root/'scan.pdf')[0]['text'].lower()
    print('OCR PASS scanned PDF',flush=True)
    (root/'sample.rtf').write_text(r'{\rtf1\ansi '+line+'}')
    assert 'sampling' in rag.extract_text_from_file(root/'sample.rtf')[0]['text']
    print('RTF PASS',flush=True)
    for ext in ['odt','ods','odp']:
        with zipfile.ZipFile(root/('sample.'+ext),'w') as z:
            z.writestr('content.xml','<document><p>'+line+'</p></document>')
        assert 'sampling' in rag.extract_text_from_file(root/('sample.'+ext))[0]['text']
        print('OPENDOCUMENT PASS',ext,flush=True)
    from docx import Document
    d=Document();d.add_paragraph(line);d.save(root/'legacy.docx')
    from pptx import Presentation
    p=Presentation();slide=p.slides.add_slide(p.slide_layouts[1]);slide.shapes.title.text=line;p.save(root/'slides.pptx')
    from openpyxl import Workbook
    w=Workbook();w.active.append([line]);w.save(root/'sheet.xlsx');w.close()
    office='C:/Program Files/LibreOffice/program/soffice.exe'
    for source,ext in ([('legacy.docx','doc'),('slides.pptx','ppt'),('sheet.xlsx','xls')] if Path(office).is_file() else []):
        subprocess.run([office,'--headless','--convert-to',ext,'--outdir',str(root),str(root/source)],check=True,timeout=120)
        assert 'sampling' in ' '.join(x['text'] for x in rag.extract_text_from_file(root/(Path(source).stem+'.'+ext))).lower()
        print('LEGACY PASS',ext,flush=True)
    script=root/'speech.ps1'
    wav=root/'speech.wav'
    script.write_text("Add-Type -AssemblyName System.Speech\n$speaker = New-Object System.Speech.Synthesis.SpeechSynthesizer\n$speaker.SetOutputToWaveFile('"+str(wav).replace("'","''")+"')\n$speaker.Speak('Stratified sampling improves survey precision. Each group is sampled separately.')\n$speaker.Dispose()")
    subprocess.run(['powershell.exe','-NoProfile','-File',str(script)],check=True,timeout=60)
    transcribed=rag.extract_text_from_file(wav)
    assert 'sampling' in ' '.join(x['text'] for x in transcribed).lower(),transcribed
    print('AUDIO TRANSCRIPTION PASS',transcribed,flush=True)
    assert Path(office).is_file(), 'LibreOffice is not yet installed; legacy DOC/PPT/XLS conversion checks are pending.'
