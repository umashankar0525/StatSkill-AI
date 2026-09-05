"""Real local AI, extraction and retrieval checks using an isolated database."""
import importlib.util
import json
import os
import sys
import tempfile
import time
from pathlib import Path
root=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(root))
tmp=tempfile.TemporaryDirectory()
os.environ['STATSKILL_DB']=str(Path(tmp.name)/'rag.db')
import storage
import rag_engine as rag
storage.init()
rag.init_vector_db()
text='''A simple random sample gives every population unit an equal chance of selection. Stratified sampling divides a population into non-overlapping groups and draws a sample from each group. Stratification can improve precision when units within each stratum are similar. A sampling frame is the list of units from which a sample is selected. Survey weights are the inverse of selection probabilities. A larger probability of selection leads to a smaller survey weight.'''
base=Path(tmp.name)
(base/'sample.txt').write_text(text)
from docx import Document
d=Document();d.add_paragraph(text);d.save(base/'sample.docx')
from pptx import Presentation
p=Presentation();s=p.slides.add_slide(p.slide_layouts[1]);s.shapes.title.text='Sampling';s.placeholders[1].text=text;p.save(base/'sample.pptx')
from openpyxl import Workbook
w=Workbook();w.active.append(['Topic','Material']);w.active.append(['Sampling',text]);w.save(base/'sample.xlsx');w.close()
import fitz
pdf=fitz.open();page=pdf.new_page();page.insert_textbox(fitz.Rect(30,30,560,750),text);pdf.save(base/'sample.pdf');pdf.close()
for ext in ('txt','docx','pptx','xlsx','pdf'):
    assert 'sampling' in ' '.join(x['text'] for x in rag.extract_text_from_file(base/('sample.'+ext))).lower()
    print('EXTRACTION PASS',ext,flush=True)
for ext in ('csv','tsv','json','md','html','xml','srt','vtt'):
    (base/('sample.'+ext)).write_text(text)
    assert rag.extract_text_from_file(base/('sample.'+ext))
    print('EXTRACTION PASS',ext,flush=True)
print('OPTIONAL DEPENDENCIES',json.dumps({m:bool(importlib.util.find_spec(m)) for m in ['faster_whisper','striprtf']}),flush=True)
print('Loading real cached embeddings…',flush=True)
start=time.time()
rag.process_and_store_document('sample','Sampling handbook','sample.txt','test',base/'sample.txt')
sources=rag.retrieve_context('What is stratified sampling?',doc_id='sample')
assert sources and sources[0]['doc_id']=='sample'
print('VECTOR RETRIEVAL PASS',round(time.time()-start,1),'seconds',flush=True)
print('Generating real source-grounded Llama question…',flush=True)
q=rag.generate_mcq_with_llm(text,'Survey sampling','Apply','Statistical Officer',{'designation':'Statistical Officer','responsibilities':'Design household surveys'})
print('REAL LLAMA PASS',json.dumps(q,ensure_ascii=False),flush=True)
