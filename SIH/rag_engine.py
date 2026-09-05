"""Local extraction, sqlite-vec retrieval and validated Ollama generation."""
import json
import os
import secrets
import shutil
import sqlite3
import threading
import urllib.request
import time
import generation_runtime as generation
from pathlib import Path
from storage import DB_PATH, db

MODEL = os.environ.get('OLLAMA_MODEL', 'llama3.1:8b')
QUIZ_MODEL = os.environ.get('STATSKILL_QUIZ_MODEL', MODEL)
HOST = os.environ.get('OLLAMA_HOST', 'http://127.0.0.1:11434').rstrip('/')
if not HOST.startswith(('http://','https://')):
    HOST = 'http://' + HOST
SUPPORTED = ['pdf','docx','pptx','xlsx','xls','csv','tsv','txt','md','json','html','xml','png','jpg','jpeg','tif','tiff','bmp','webp','srt','vtt','odt','ods','odp','rtf','doc','ppt','mp3','wav','m4a','mp4','webm','mov']
_embedder = None
_embed_lock = threading.RLock()
_llm_lock = threading.Lock()

def get_embedder():
    global _embedder
    with _embed_lock:
        if _embedder is None:
            from sentence_transformers import SentenceTransformer
            _embedder = SentenceTransformer(os.environ.get('STATSKILL_EMBED_MODEL','all-MiniLM-L6-v2'), local_files_only=True)
        return _embedder

def vector_connection():
    import sqlite_vec
    c = sqlite3.connect(DB_PATH, timeout=30)
    c.row_factory = sqlite3.Row
    c.enable_load_extension(True)
    sqlite_vec.load(c)
    c.enable_load_extension(False)
    return c

def init_vector_db():
    c = vector_connection()
    try:
        c.executescript('''CREATE TABLE IF NOT EXISTS documents(id TEXT PRIMARY KEY,title TEXT,filename TEXT,uploaded_by TEXT,created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS doc_chunks(rowid INTEGER PRIMARY KEY AUTOINCREMENT,doc_id TEXT,chunk_text TEXT,page_num INTEGER);
        CREATE VIRTUAL TABLE IF NOT EXISTS doc_chunks_vec USING vec0(chunk_embedding float[384]);''')
        cols = {r['name'] for r in c.execute('PRAGMA table_info(documents)')}
        for name, definition in [('status',"TEXT NOT NULL DEFAULT 'pending'"),('error','TEXT'),('storage_path','TEXT')]:
            if name not in cols:
                c.execute(f'ALTER TABLE documents ADD COLUMN {name} {definition}')
        c.commit()
    finally:
        c.close()

def ocr(image):
    import pytesseract
    configured = os.environ.get('TESSERACT_CMD') or shutil.which('tesseract')
    if not configured:
        for candidate in [Path(os.environ.get('LOCALAPPDATA',''))/'Programs/Tesseract-OCR/tesseract.exe',Path('C:/Program Files/Tesseract-OCR/tesseract.exe')]:
            if candidate.is_file():
                configured = str(candidate)
                break
    if not configured:
        raise ValueError('OCR requires Tesseract. Install it and set TESSERACT_CMD, or upload a text-based document.')
    pytesseract.pytesseract.tesseract_cmd = configured
    return pytesseract.image_to_string(image, timeout=120)

def extract_text_from_file(file_path):
    p = Path(file_path)
    ext = p.suffix.lower()[1:]
    pages = []
    if ext == 'pdf':
        import fitz
        from PIL import Image
        with fitz.open(p) as doc:
            for n,page in enumerate(doc):
                content = page.get_text()
                if not content.strip():
                    pix = page.get_pixmap(matrix=fitz.Matrix(2,2))
                    content = ocr(Image.frombytes('RGB',[pix.width,pix.height],pix.samples))
                pages.append({'text':content,'page':n+1})
    elif ext == 'docx':
        from docx import Document
        d = Document(p)
        content = '\n'.join(x.text for x in d.paragraphs)
        content += '\n' + '\n'.join(' | '.join(cell.text for cell in row.cells) for t in d.tables for row in t.rows)
        pages = [{'text':content,'page':1}]
    elif ext == 'pptx':
        from pptx import Presentation
        for n,slide in enumerate(Presentation(p).slides):
            parts = [s.text for s in slide.shapes if s.has_text_frame]
            for s in slide.shapes:
                if s.has_table:
                    parts.extend(' | '.join(cell.text for cell in r.cells) for r in s.table.rows)
            if slide.has_notes_slide:
                parts.append(slide.notes_slide.notes_text_frame.text)
            pages.append({'text':'\n'.join(parts),'page':n+1})
    elif ext in ('xlsx','xls'):
        if ext == 'xlsx':
            import openpyxl
            workbook = openpyxl.load_workbook(p,read_only=True,data_only=True)
            try:
                for n,sheet in enumerate(workbook):
                    pages.append({'text':sheet.title+'\n'+'\n'.join(' | '.join(str(v) if v is not None else '' for v in row) for row in sheet.values),'page':n+1})
            finally:
                workbook.close()
        else:
            import xlrd
            workbook = xlrd.open_workbook(p)
            for n,sheet in enumerate(workbook.sheets()):
                pages.append({'text':sheet.name+'\n'+'\n'.join(' | '.join(map(str,sheet.row_values(i))) for i in range(sheet.nrows)),'page':n+1})
    elif ext in ('odt','ods','odp'):
        import zipfile
        import xml.etree.ElementTree as ET
        with zipfile.ZipFile(p) as z:
            root = ET.fromstring(z.read('content.xml'))
            pages = [{'text':'\n'.join(root.itertext()),'page':1}]
    elif ext == 'rtf':
        from striprtf.striprtf import rtf_to_text
        pages = [{'text':rtf_to_text(p.read_text(encoding='utf-8',errors='replace')),'page':1}]
    elif ext in ('doc','ppt'):
        import tempfile
        import subprocess
        soffice = shutil.which('soffice') or os.environ.get('LIBREOFFICE_CMD')
        if not soffice and Path('C:/Program Files/LibreOffice/program/soffice.exe').is_file():
            soffice='C:/Program Files/LibreOffice/program/soffice.exe'
        if not soffice:
            raise ValueError('Legacy DOC/PPT needs LibreOffice. Save as DOCX/PPTX or configure LIBREOFFICE_CMD.')
        with tempfile.TemporaryDirectory() as tmp:
            fmt = 'docx' if ext == 'doc' else 'pptx'
            subprocess.run([soffice,'--headless','--convert-to',fmt,'--outdir',tmp,str(p)],check=True,timeout=120)
            return extract_text_from_file(Path(tmp)/(p.stem+'.'+fmt))
    elif ext in ('mp3','wav','m4a','mp4','webm','mov'):
        from faster_whisper import WhisperModel
        model = WhisperModel(os.environ.get('WHISPER_MODEL','base'),device='cpu',compute_type='int8',local_files_only=True)
        segments,_ = model.transcribe(str(p))
        pages = [{'text':f'[{s.start:.1f}s–{s.end:.1f}s] {s.text}','page':i+1} for i,s in enumerate(segments)]
    elif ext in ('png','jpg','jpeg','tif','tiff','bmp','webp'):
        from PIL import Image,ImageSequence
        with Image.open(p) as im:
            pages = [{'text':ocr(frame.convert('RGB')),'page':i+1} for i,frame in enumerate(ImageSequence.Iterator(im))]
    elif ext in ('txt','md','csv','tsv','json','html','xml','srt','vtt'):
        import chardet
        raw = p.read_bytes()
        content = raw.decode(chardet.detect(raw).get('encoding') or 'utf-8')
        if ext == 'html':
            from html.parser import HTMLParser
            class TextParser(HTMLParser):
                def __init__(self):
                    super().__init__()
                    self.parts=[]
                    self.hidden=0
                def handle_starttag(self,tag,attrs):
                    if tag in ('script','style'):
                        self.hidden+=1
                def handle_endtag(self,tag):
                    if tag in ('script','style'):
                        self.hidden=max(0,self.hidden-1)
                def handle_data(self,data):
                    if not self.hidden:
                        self.parts.append(data)
            parser=TextParser()
            parser.feed(content)
            content=' '.join(parser.parts)
        pages = [{'text':content,'page':1}]
    else:
        raise ValueError('Unsupported file format. See the supported format list.')
    if not any(x['text'].strip() for x in pages):
        raise ValueError('No readable text found. Upload a document with text or enable OCR.')
    if sum(len(x['text']) for x in pages) > 2_000_000:
        raise ValueError('Extracted text exceeds 2 million characters. Split the material into smaller files.')
    return pages

def process_and_store_document(doc_id, title, filename, uploaded_by, file_path):
    pages = extract_text_from_file(file_path)
    return process_and_store_pages(doc_id,title,filename,uploaded_by,pages)

def process_and_store_pages(doc_id,title,filename,uploaded_by,pages):
    chunks = [{'text':p['text'][i:i+1000], 'page':p['page']} for p in pages for i in range(0,len(p['text']),800) if p['text'][i:i+1000].strip()]
    with _embed_lock:
        vectors = get_embedder().encode([x['text'] for x in chunks],batch_size=32).tolist()
    c = vector_connection()
    try:
        c.execute('INSERT OR IGNORE INTO documents(id,title,filename,uploaded_by,status) VALUES(?,?,?,?,?)',(doc_id,title,filename,uploaded_by,'pending'))
        for chunk,vector in zip(chunks,vectors):
            row = c.execute('INSERT INTO doc_chunks(doc_id,chunk_text,page_num) VALUES(?,?,?)',(doc_id,chunk['text'],chunk['page']))
            c.execute('INSERT INTO doc_chunks_vec(rowid,chunk_embedding) VALUES(?,?)',(row.lastrowid,json.dumps(vector)))
        c.execute("UPDATE documents SET status='ready',error=NULL WHERE id=?",(doc_id,))
        c.commit()
    except Exception:
        c.rollback()
        raise
    finally:
        c.close()
    return len(chunks)

def retrieve_context(query, top_k=3, doc_id=None, doc_ids=None):
    if doc_ids is not None and not doc_ids:return []
    c = vector_connection()
    try:
        if not c.execute("SELECT 1 FROM documents WHERE status='ready' LIMIT 1").fetchone():
            return []
        with _embed_lock:
            embedding = json.dumps(get_embedder().encode(query).tolist())
        # Filter BEFORE ranking: global KNN then a document filter loses relevant hits.
        scope=''
        params=[embedding,doc_id,doc_id]
        if doc_ids is not None:
            scope=' AND c.doc_id IN ('+','.join('?' for _ in doc_ids)+')'
            params.extend(doc_ids)
        params.append(top_k)
        rows = c.execute('''SELECT c.rowid,c.doc_id,c.chunk_text,c.page_num,d.title,
          vec_distance_cosine(v.chunk_embedding,?) AS distance FROM doc_chunks c
          JOIN documents d ON d.id=c.doc_id JOIN doc_chunks_vec v ON v.rowid=c.rowid
          WHERE d.status='ready' AND (? IS NULL OR c.doc_id=?)'''+scope+' ORDER BY distance LIMIT ?',params).fetchall()
        return [dict(id=r['rowid'],doc_id=r['doc_id'],text=r['chunk_text'],page=r['page_num'],title=r['title'],score=r['distance']) for r in rows]
    finally:
        c.close()

def chat(messages, json_output=False, model=None, max_tokens=650, temperature=0.2, context_length=None, gpu_layers=None):
    payload = {'model':model or MODEL,'messages':messages,'stream':False,'keep_alive':'30m',
               'options':{'num_ctx':int(context_length or os.environ.get('STATSKILL_CONTEXT_LENGTH','8192')),'num_predict':max_tokens,'temperature':temperature,
                          'num_gpu':int(gpu_layers if gpu_layers is not None else os.environ.get('STATSKILL_GPU_LAYERS','-1'))}}
    if json_output:
        payload['format'] = json_output if isinstance(json_output,dict) else 'json'
    started=time.monotonic()
    with _llm_lock:
        queued=time.monotonic()-started
        generation.event('request_started',model=payload['model'],prompt_chars=sum(len(m['content']) for m in messages),token_limit=max_tokens,context_length=payload['options']['num_ctx'],gpu_layers=payload['options']['num_gpu'],queue_seconds=round(queued,3))
        request = urllib.request.Request(HOST+'/api/chat',data=json.dumps(payload).encode(),headers={'Content-Type':'application/json'})
        try:
            with urllib.request.urlopen(request,timeout=float(os.environ.get('STATSKILL_MODEL_TIMEOUT','300'))) as r:
                data = json.load(r)
            generation.event('request_finished',model=payload['model'],elapsed_seconds=round(time.monotonic()-started-queued,3),done_reason=data.get('done_reason'),prompt_tokens=data.get('prompt_eval_count'),output_tokens=data.get('eval_count'),load_seconds=data.get('load_duration',0)/1e9)
            if not data.get('done') or data.get('done_reason') == 'length':
                raise generation.GenerationError('truncated','The local model reached its response limit.')
            return data['message']['content']
        except Exception as e:
            failure=generation.classify(e)
            generation.event('request_failed',model=payload['model'],error_code=failure.code,elapsed_seconds=round(time.monotonic()-started,3))
            raise failure from e

def generate_mcq_with_llm(context_text, competency_name, bloom_level_name, role_title, profile=None, previous=(), model_override=None):
    from soul_quiz_engine import generate_mcq
    return generate_mcq(context_text,competency_name,bloom_level_name,role_title,profile,previous,model_override)

def generate_mcqs_with_llm(context_text,competency_name,bloom_level_name,role_title,profile=None,previous=(),count=2):
    from soul_quiz_engine import generate_batch
    return generate_batch(context_text,competency_name,bloom_level_name,role_title,profile,previous,count)
